"""Render both patent applications as DOCX — the NIIS portal accepts only doc/docx.

Same sources as make_pdf.py (um/, inv/ Markdown + figures/), same layout rules:
A4, margins 25/20/20/20 mm, Times New Roman 12 pt, 1.5 line spacing, first-line
indent 1 cm, bullet lists, Markdown tables, one drawing per page captioned «Фиг. N».

Usage (from the repo root):
    python ip/patent/make_docx.py            # both applications -> build/<app>/*.docx
    python ip/patent/make_docx.py --only um
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

BASE = Path(__file__).resolve().parent
BUILD = BASE / "build"
FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))")


def _new_document() -> Document:
    """A4 document with the patent margins and a Times 12 pt / 1.5 base style."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(25), Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(0)
    return doc


def _add_runs(par, text: str) -> None:
    """Append text with the Markdown inline subset (**bold**, *italic*, `code`)."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            par.add_run(part[1:-1])
        elif part.startswith("*") and part.endswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _para(doc: Document, text: str, *, indent: bool = True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    par = doc.add_paragraph()
    par.alignment = align
    if indent:
        par.paragraph_format.first_line_indent = Cm(1)
    _add_runs(par, text)
    return par


def _table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            p = table.cell(i, j).paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            _add_runs(p, cell)
            for run in p.runs:
                run.font.size = Pt(10)
                run.bold = run.bold or i == 0
    doc.add_paragraph()


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def render_markdown(source: Path, target: Path) -> None:
    """Convert one application document (description, claims or abstract) to DOCX."""
    doc = _new_document()
    para: list[str] = []
    bullet: list[str] = []
    rows: list[list[str]] = []

    def flush() -> None:
        nonlocal para, bullet, rows
        if para:
            _para(doc, " ".join(para))
        if bullet:
            p = _para(doc, "– " + " ".join(bullet), indent=False)
            p.paragraph_format.left_indent = Cm(1)
        if rows:
            _table(doc, rows)
        para, bullet, rows = [], [], []

    for raw in source.read_text(encoding="utf-8").splitlines():
        s = raw.strip()
        if s.startswith("<!--"):
            continue
        if s.startswith("|"):
            if para or bullet:
                flush()
            if not set(s.replace("|", "").strip()) <= set("-: "):
                rows.append(_cells(s))
            continue
        if rows:
            flush()
        if not s:
            flush()
        elif s.startswith("# "):
            flush()
            p = _para(doc, s[2:], indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
            for r in p.runs:
                r.bold = True
        elif s.startswith("## "):
            flush()
            p = _para(doc, s[3:], indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_before = Pt(6)
            for r in p.runs:
                r.bold = True
        elif s.startswith("МПК"):
            flush()
            _para(doc, s, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif s.startswith("- "):
            flush()
            bullet = [s[2:]]
        elif re.match(r"\d+\. ", s):
            flush()
            para = [s]
        elif bullet:
            bullet.append(s)
        else:
            para.append(s)
    flush()
    doc.save(target)


def render_figures(fig_dir: Path, target: Path) -> None:
    """One drawing per page, scaled to the text width, captioned «Фиг. N»."""
    doc = _new_document()
    figs = sorted(fig_dir.glob("fig*.png"), key=lambda p: int(p.stem[3:]))
    for i, path in enumerate(figs):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Mm(165))
        cap = doc.add_paragraph(f"Фиг. {path.stem[3:]}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if i < len(figs) - 1:
            cap.add_run().add_break(WD_BREAK.PAGE)
    doc.save(target)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", choices=["um", "inv"])
    args = ap.parse_args()
    for app in ([args.only] if args.only else ["um", "inv"]):
        src, out = BASE / app, BUILD / app
        out.mkdir(parents=True, exist_ok=True)
        for name in ("opisanie", "formula", "referat"):
            render_markdown(src / f"{name}.md", out / f"{name}.docx")
        render_figures(src / "figures", out / "figures.docx")
        print(f"{app}: " + ", ".join(p.name for p in sorted(out.glob("*.docx"))))


if __name__ == "__main__":
    main()
